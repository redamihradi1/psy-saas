"""Export d'un bilan en Word (.docx, python-docx) et en PDF (reportlab), à partir des mêmes
blocs que l'éditeur (html_blocs). Mise en page d'un bilan clinique : papier à en-tête du
psychologue (EnTeteBilan), page de garde, titres bleu marine, tableaux à en-tête foncé et lignes
alternées, graphiques de profil, sauts de page avant les grandes parties, signature."""
import io
import re
from xml.sax.saxutils import escape

from django.utils.text import slugify

from . import graphiques
from . import html_blocs as hb
from . import tests_externes as te
from .modeles import section_du_modele

MARINE = (0x1F, 0x3A, 0x5F)
ENTETE_TABLEAU = (0x24, 0x48, 0x6F)
LIGNE_ALTERNEE = (0xEE, 0xF3, 0xF9)
BORDURE = (0xC9, 0xD6, 0xE3)
GRIS = (0x55, 0x55, 0x55)
RE_ARABE = re.compile(r'[؀-ۿ]')
SECTIONS_SANS_TITRE = {'identite', 'signature'}


def nom_fichier(bilan, extension):
    base = slugify(f"bilan-{bilan.patient.nom}-{bilan.patient.prenom}-{bilan.date_modification:%Y%m%d}")
    return f"{base or 'bilan'}.{extension}"


def _hex(rgb):
    return '%02X%02X%02X' % rgb


# --- éléments communs ---------------------------------------------------------------------------

def _en_tete(bilan):
    from .donnees import en_tete_de
    return en_tete_de(bilan)


def _lignes(texte):
    return [l.strip() for l in (texte or '').splitlines() if l.strip()]


def _bloc_identite_psy(bilan, en_tete):
    """-> (lignes de gauche [(texte, style)], lignes de droite, lignes du pied)."""
    org = bilan.organization
    if en_tete:
        gauche = [(en_tete.nom_affiche, 'nom'), (en_tete.titre_pro, 'titre'), (en_tete.specialites, 'italique')]
        gauche = [(t, s) for t, s in gauche if t]
        droite = _lignes(en_tete.texte_droite)
        pied = ([en_tete.mention_confidentielle] if en_tete.mention_confidentielle else []) + _lignes(en_tete.pied_de_page)
    else:
        gauche = [(org.name, 'nom')]
        coordonnees = ' - '.join(x for x in [org.address.replace('\n', ', ') if org.address else '', org.city, org.phone] if x)
        if coordonnees:
            gauche.append((coordonnees, 'titre'))
        droite = []
        pied = ["Document clinique confidentiel"]
    return gauche, droite, pied


def png_figure(bilan, ref):
    """Image d'un bloc figure : graphique de profil (« wisc5.indices ») ou signature."""
    if ref == 'signature':
        en_tete = _en_tete(bilan)
        return graphiques.image_png(en_tete.signature) if en_tete and en_tete.signature else None
    cle, _, partie = ref.partition('.')
    saisie = (bilan.scores or {}).get(cle)
    if saisie is None:
        return None
    donnees_fig = te.donnees_figure(cle, saisie, partie)
    return graphiques.profil_png(donnees_fig) if donnees_fig else None


def _est_legende(bloc, suivant):
    """Paragraphe tout en gras juste avant un tableau (« Indices », « Subtests ») : il reste sur la
    même page que son tableau."""
    runs = [r for r in bloc.get('runs', []) if 't' in r]
    return bloc['type'] == 'p' and suivant is not None and suivant['type'] == 'table' and runs and all(r.get('b') for r in runs)


def _parcours(bilan):
    """Blocs du document annotés pour la mise en page :
    -> [(bloc, {'section': clé, 'saut': bool, 'masquer': bool, 'legende': bool})]"""
    resultat = []
    blocs = hb.lire(bilan.contenu)
    section = ''
    premier_test = True
    for n, bloc in enumerate(blocs):
        info = {'section': section, 'saut': False, 'masquer': False,
                'legende': _est_legende(bloc, blocs[n + 1] if n + 1 < len(blocs) else None)}
        if bloc['type'] == 'h2' and bloc.get('section'):
            section = bloc['section']
            info['section'] = section
            definition = section_du_modele(bilan.modele, section)
            if bloc.get('section_type') == 'test':
                info['saut'] = premier_test
                premier_test = False
            elif definition and definition.get('page'):
                info['saut'] = True
            info['masquer'] = section in SECTIONS_SANS_TITRE
        resultat.append((bloc, info))
    return resultat


# --- Word ------------------------------------------------------------------------------------

def exporter_docx(bilan):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    marine = RGBColor(*MARINE)
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = Cm(3.2)
    section.header_distance = Cm(0.8)
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    style.font.size = Pt(11.5)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.25

    def ombrer(cellule, rgb):
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), _hex(rgb))
        cellule._tc.get_or_add_tcPr().append(shd)

    def bordure_basse(paragraphe, rgb=MARINE, epaisseur=6):
        pPr = paragraphe._p.get_or_add_pPr()
        bordures = OxmlElement('w:pBdr')
        bas = OxmlElement('w:bottom')
        for attr, valeur in (('w:val', 'single'), ('w:sz', str(epaisseur)), ('w:space', '2'), ('w:color', _hex(rgb))):
            bas.set(qn(attr), valeur)
        bordures.append(bas)
        pPr.append(bordures)

    def bordures_tableau(table, rgb=BORDURE):
        tblPr = table._tbl.tblPr
        bordures = OxmlElement('w:tblBorders')
        for cote in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            b = OxmlElement(f'w:{cote}')
            for attr, valeur in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '0'), ('w:color', _hex(rgb))):
                b.set(qn(attr), valeur)
            bordures.append(b)
        tblPr.append(bordures)

    def texte_arabe(run):
        rPr = run._element.get_or_add_rPr()
        rPr.append(OxmlElement('w:rtl'))

    def champ_page(paragraphe):
        run = paragraphe.add_run()
        for type_champ, texte in (('begin', None), (None, 'PAGE'), ('end', None)):
            if type_champ:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), type_champ)
            else:
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = texte
            run._r.append(el)
        run.font.size = Pt(8)

    # En-tête : identité du psychologue à gauche, logo au centre, texte libre à droite
    en_tete = _en_tete(bilan)
    gauche, droite, pied = _bloc_identite_psy(bilan, en_tete)
    logo = graphiques.image_png(en_tete.logo) if en_tete and en_tete.logo else None
    table_entete = section.header.add_table(1, 3, Cm(16.6))
    table_entete.alignment = WD_TABLE_ALIGNMENT.CENTER
    cellules = table_entete.rows[0].cells
    for n, (texte, genre) in enumerate(gauche):
        p = cellules[0].paragraphs[0] if n == 0 else cellules[0].add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(texte)
        r.font.size = Pt(12 if genre == 'nom' else 10)
        r.italic = genre == 'italique'
        if genre == 'nom':
            r.font.color.rgb = marine
    if logo:
        p = cellules[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(io.BytesIO(logo), height=Cm(1.9))
    for n, ligne in enumerate(droite):
        p = cellules[2].paragraphs[0] if n == 0 else cellules[2].add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(ligne)
        r.font.size = Pt(11 if n == 0 else 10)
        if n == 0:
            r.font.color.rgb = marine
        if RE_ARABE.search(ligne):
            texte_arabe(r)

    # Pied de page : mention confidentielle, coordonnées, numéro de page
    pied_p = section.footer.paragraphs[0]
    pied_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for n, ligne in enumerate(pied):
        if n:
            pied_p.add_run().add_break()
        r = pied_p.add_run(ligne)
        r.font.size = Pt(8)
        r.bold = n == 0
        if n:
            r.font.color.rgb = marine
    page = section.footer.add_paragraph()
    page.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    champ_page(page)

    def ajouter_runs(paragraphe, runs, taille=None):
        for run in runs:
            if 'br' in run:
                paragraphe.add_run().add_break()
                continue
            r = paragraphe.add_run(run['t'])
            r.bold, r.italic, r.underline = run.get('b'), run.get('i'), run.get('u')
            if taille:
                r.font.size = Pt(taille)

    def saut_de_page():
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def tableau(bloc, identite=False):
        nb_colonnes = max(len(row['cells']) for row in bloc['rows'])
        table = document.add_table(rows=0, cols=nb_colonnes)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordures_tableau(table)
        n_corps = 0
        for row in bloc['rows']:
            cellules = table.add_row().cells
            entete = row.get('header')
            for n, cellule_runs in enumerate(row['cells']):
                p = cellules[n].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                ajouter_runs(p, cellule_runs, taille=10 if identite else 9.5)
                if identite:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if n == 0 else WD_ALIGN_PARAGRAPH.CENTER
                elif n > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if entete:
                    ombrer(cellules[n], ENTETE_TABLEAU)
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                elif n_corps % 2 == 1 and (not identite or n == 1):
                    ombrer(cellules[n], LIGNE_ALTERNEE)
            if not entete:
                n_corps += 1
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    debut_document = True
    for bloc, info in _parcours(bilan):
        t = bloc['type']
        if info['saut'] and not debut_document:
            saut_de_page()
        if t == 'h1':
            for _ in range(5):
                document.add_paragraph()
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            ajouter_runs(p, bloc['runs'], taille=26)
            for r in p.runs:
                r.bold, r.font.color.rgb = True, marine
            bordure_basse(p, epaisseur=8)
        elif t == 'h2':
            if info['masquer']:
                continue
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.keep_with_next = True
            ajouter_runs(p, bloc['runs'], taille=14)
            for r in p.runs:
                r.bold, r.font.color.rgb = True, marine
            bordure_basse(p, rgb=BORDURE)
        elif t == 'h3':
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            ajouter_runs(p, bloc['runs'], taille=12)
            for r in p.runs:
                r.bold, r.font.color.rgb = True, marine
        elif t == 'p':
            p = document.add_paragraph()
            if info['section'] == '' and not debut_document:  # sous-titre de la page de garde
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                ajouter_runs(p, bloc['runs'], taille=12)
                for r in p.runs:
                    r.font.color.rgb = marine
            elif info['section'] == 'signature':
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                ajouter_runs(p, bloc['runs'])
            elif info['legende']:
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(6)
                ajouter_runs(p, bloc['runs'])
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                ajouter_runs(p, bloc['runs'])
        elif t in ('ul', 'ol'):
            for item in bloc['items']:
                p = document.add_paragraph(style='List Bullet' if t == 'ul' else 'List Number')
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                ajouter_runs(p, item)
        elif t == 'table':
            tableau(bloc, identite=info['section'] == 'identite')
        elif t == 'figure':
            png = png_figure(bilan, bloc['ref'])
            if png:
                p = document.add_paragraph()
                signature = bloc['ref'] == 'signature'
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if signature else WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(io.BytesIO(png), width=Cm(5) if signature else Cm(16))
        debut_document = False

    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


# --- PDF -------------------------------------------------------------------------------------

# Polices standard PDF (Times) = encodage WinAnsi : on remplace les quelques caractères hors de
# cet encodage que l'IA ou le psy peuvent taper.
REMPLACEMENTS_PDF = {'≥': '>=', '≤': '<=', '→': '->', '←': '<-', '≠': '!=', '−': '-', ' ': ' ', ' ': ' '}


def _texte_pdf(texte):
    for a, b in REMPLACEMENTS_PDF.items():
        texte = texte.replace(a, b)
    return texte.encode('cp1252', 'replace').decode('cp1252')


def _affichable_pdf(texte):
    """Les lignes en arabe (non couvertes par les polices standard) ne vont que dans le Word."""
    return not RE_ARABE.search(texte or '')


def _markup(runs):
    morceaux = []
    for run in runs:
        if 'br' in run:
            morceaux.append('<br/>')
            continue
        texte = escape(_texte_pdf(run['t']))
        if run.get('u'):
            texte = f'<u>{texte}</u>'
        if run.get('i'):
            texte = f'<i>{texte}</i>'
        if run.get('b'):
            texte = f'<b>{texte}</b>'
        morceaux.append(texte)
    return ''.join(morceaux)


def exporter_pdf(bilan):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (HRFlowable, Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    def couleur(rgb):
        return colors.Color(*(c / 255 for c in rgb))

    marine, entete_tab, alternee, bordure = couleur(MARINE), couleur(ENTETE_TABLEAU), couleur(LIGNE_ALTERNEE), couleur(BORDURE)
    base = getSampleStyleSheet()
    styles = {
        'h1': ParagraphStyle('h1', parent=base['Title'], fontName='Times-Bold', fontSize=24, leading=28,
                             textColor=marine, alignment=TA_RIGHT, spaceAfter=4),
        'sous_titre': ParagraphStyle('st', parent=base['Normal'], fontName='Times-Roman', fontSize=12, leading=15,
                                     textColor=marine, alignment=TA_RIGHT, spaceAfter=18),
        'h2': ParagraphStyle('h2', parent=base['Heading2'], fontName='Times-Bold', fontSize=14, leading=17,
                             textColor=marine, spaceBefore=12, spaceAfter=2),
        'h3': ParagraphStyle('h3', parent=base['Heading3'], fontName='Times-Bold', fontSize=12, leading=15,
                             textColor=marine, spaceBefore=8, spaceAfter=4),
        'p': ParagraphStyle('p', parent=base['BodyText'], fontName='Times-Roman', fontSize=11.5, leading=16,
                            alignment=TA_JUSTIFY, spaceAfter=6),
        'legende': ParagraphStyle('lg', parent=base['BodyText'], fontName='Times-Bold', fontSize=11.5, leading=15,
                                  spaceBefore=6, spaceAfter=4, keepWithNext=1),
        'droite': ParagraphStyle('dr', parent=base['BodyText'], fontName='Times-Roman', fontSize=11.5, leading=15,
                                 alignment=TA_RIGHT, spaceAfter=3),
        'cellule': ParagraphStyle('c', parent=base['BodyText'], fontName='Times-Roman', fontSize=9.5, leading=11.5),
        'cellule_centre': ParagraphStyle('cc', parent=base['BodyText'], fontName='Times-Roman', fontSize=9.5,
                                         leading=11.5, alignment=TA_CENTER),
        'cellule_entete': ParagraphStyle('ce', parent=base['BodyText'], fontName='Times-Bold', fontSize=9.5,
                                         leading=11.5, textColor=colors.white, alignment=TA_CENTER),
        'identite_libelle': ParagraphStyle('il', parent=base['BodyText'], fontName='Times-Italic', fontSize=10.5,
                                           leading=13, alignment=TA_RIGHT),
        'identite_valeur': ParagraphStyle('iv', parent=base['BodyText'], fontName='Times-Roman', fontSize=10.5,
                                          leading=13, alignment=TA_CENTER),
    }

    en_tete = _en_tete(bilan)
    gauche, droite, pied = _bloc_identite_psy(bilan, en_tete)
    gauche = [(_texte_pdf(t), g) for t, g in gauche]
    droite = [_texte_pdf(l) for l in droite if _affichable_pdf(l)]
    pied = [_texte_pdf(l) for l in pied if _affichable_pdf(l)]
    logo_png = graphiques.image_png(en_tete.logo) if en_tete and en_tete.logo else None
    logo = ImageReader(io.BytesIO(logo_png)) if logo_png else None

    def entete_pied(canvas, doc):
        canvas.saveState()
        y = A4[1] - 1.3 * cm
        for texte, genre in gauche:
            canvas.setFillColor(marine if genre == 'nom' else colors.black)
            canvas.setFont({'nom': 'Times-Roman', 'italique': 'Times-Italic'}.get(genre, 'Times-Roman'),
                           12 if genre == 'nom' else 10)
            canvas.drawString(2 * cm, y, texte)
            y -= 0.5 * cm
        y = A4[1] - 1.3 * cm
        for n, ligne in enumerate(droite):
            canvas.setFillColor(marine if n == 0 else colors.black)
            canvas.setFont('Times-Roman', 11 if n == 0 else 10)
            canvas.drawRightString(A4[0] - 2 * cm, y, ligne)
            y -= 0.5 * cm
        if logo:
            largeur, hauteur = logo.getSize()
            h = 1.9 * cm
            w = h * largeur / hauteur
            canvas.drawImage(logo, (A4[0] - w) / 2, A4[1] - 0.8 * cm - h, width=w, height=h, mask='auto')
        y = 1.5 * cm
        for n, ligne in enumerate(pied):
            canvas.setFillColor(colors.black if n == 0 else marine)
            canvas.setFont('Times-Bold' if n == 0 else 'Times-Roman', 8 if n == 0 else 7.5)
            canvas.drawCentredString(A4[0] / 2, y, ligne)
            y -= 0.32 * cm
        canvas.setFillColor(colors.grey)
        canvas.setFont('Times-Roman', 8)
        canvas.drawRightString(A4[0] - 2 * cm, 0.8 * cm, str(doc.page))
        canvas.restoreState()

    largeur = A4[0] - 4.4 * cm

    def tableau(bloc, identite=False):
        nb_colonnes = max(len(row['cells']) for row in bloc['rows'])
        donnees_tab, style = [], [('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                                  ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
        n_corps = 0
        for n, row in enumerate(bloc['rows']):
            cellules = []
            for c, cellule in enumerate(row['cells']):
                if row.get('header'):
                    st = styles['cellule_entete']
                elif identite:
                    st = styles['identite_libelle'] if c == 0 else styles['identite_valeur']
                else:
                    st = styles['cellule'] if c == 0 else styles['cellule_centre']
                cellules.append(Paragraph(_markup(cellule) or '&nbsp;', st))
            donnees_tab.append(cellules + [''] * (nb_colonnes - len(cellules)))
            if row.get('header'):
                style.append(('BACKGROUND', (0, n), (-1, n), entete_tab))
            else:
                if n_corps % 2 == 1:
                    style.append(('BACKGROUND', (1 if identite else 0, n), (-1, n), alternee))
                n_corps += 1
        if identite:
            largeurs = [largeur * 0.38, largeur * 0.62]
            style += [('LINEBELOW', (0, 0), (-1, -1), 0.4, bordure), ('BOX', (1, 0), (1, -1), 0.4, bordure)]
        else:
            largeurs = [largeur * (0.34 if nb_colonnes > 2 else 0.5)] + \
                       [largeur * (0.66 if nb_colonnes > 2 else 0.5) / (nb_colonnes - 1)] * (nb_colonnes - 1) \
                if nb_colonnes > 1 else [largeur]
            style.append(('GRID', (0, 0), (-1, -1), 0.4, bordure))
        table = Table(donnees_tab, colWidths=largeurs, repeatRows=1 if bloc['rows'][0].get('header') else 0)
        table.setStyle(TableStyle(style))
        return [table, Spacer(1, 0.3 * cm)]

    elements = []
    for bloc, info in _parcours(bilan):
        t = bloc['type']
        if info['saut'] and elements:
            elements.append(PageBreak())
        if t == 'h1':
            elements.extend([Spacer(1, 5 * cm), Paragraph(_markup(bloc['runs']), styles['h1']),
                             HRFlowable(width='100%', thickness=1, color=marine, spaceAfter=6)])
        elif t == 'h2':
            if not info['masquer']:
                elements.extend([Paragraph(_markup(bloc['runs']), styles['h2']),
                                 HRFlowable(width='100%', thickness=0.5, color=bordure, spaceAfter=6)])
        elif t == 'h3':
            elements.append(Paragraph(_markup(bloc['runs']), styles['h3']))
        elif t == 'p':
            if info['section'] == '' and elements:
                st = styles['sous_titre']
            elif info['section'] == 'signature':
                st = styles['droite']
            elif info['legende']:
                st = styles['legende']
            else:
                st = styles['p']
            elements.append(Paragraph(_markup(bloc['runs']) or '&nbsp;', st))
        elif t in ('ul', 'ol'):
            items = [ListItem(Paragraph(_markup(item), styles['p']), leftIndent=14) for item in bloc['items']]
            elements.append(ListFlowable(items, bulletType='bullet' if t == 'ul' else '1',
                                         start='•' if t == 'ul' else None, leftIndent=16, bulletFontSize=8))
        elif t == 'table':
            elements.extend(tableau(bloc, identite=info['section'] == 'identite'))
        elif t == 'figure':
            png = png_figure(bilan, bloc['ref'])
            if png:
                lecteur = ImageReader(io.BytesIO(png))
                w_px, h_px = lecteur.getSize()
                signature = bloc['ref'] == 'signature'
                w = 5 * cm if signature else largeur
                image = Image(io.BytesIO(png), width=w, height=w * h_px / w_px)
                image.hAlign = 'RIGHT' if signature else 'CENTER'
                elements.append(KeepTogether([image, Spacer(1, 0.3 * cm)]))

    tampon = io.BytesIO()
    marge_haut = 3.4 * cm if (gauche or droite or logo) else 2.4 * cm
    doc = SimpleDocTemplate(tampon, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=marge_haut, bottomMargin=2.4 * cm, title=_texte_pdf(bilan.titre))
    doc.build(elements or [Spacer(1, 1)], onFirstPage=entete_pied, onLaterPages=entete_pied)
    return tampon.getvalue()
